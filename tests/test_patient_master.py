"""主表边界测试使用合成数据，不读取真实EEG。"""

import csv
import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

SCRIPT_PATH = (
    Path(__file__).resolve().parents[1]
    / "preprocessing/01_build_patient_master.py"
)
specification = importlib.util.spec_from_file_location(
    "patient_master", SCRIPT_PATH,
)
master_module = importlib.util.module_from_spec(specification)
sys.modules[specification.name] = master_module
specification.loader.exec_module(master_module)

Assessment = master_module.Assessment
SCORE_COLUMNS = master_module.SCORE_COLUMNS
choose_edf_path = master_module.choose_edf_path
choose_guid = master_module.choose_guid
review_scores = master_module.review_scores
parse_score = master_module.parse_score
read_assessments = master_module.read_assessments

TEST_GUID = "e2ee5b31-758c-46d4-86d8-0b7f0dcb5cd9"
OTHER_GUID = "22c8b7dd-b8ad-4fed-9891-10066c27f639"


class MasterTableTests(unittest.TestCase):
    def make_assessment(self, description: str) -> Assessment:
        return Assessment(
            "303-MWS", "MWS", "测试患者",
            dict.fromkeys(SCORE_COLUMNS), description, "测试表", 2,
        )

    def test_missing_score_is_not_zero(self) -> None:
        self.assertIsNone(parse_score("—", "测试/听觉"))
        self.assertEqual(parse_score("0", "测试/听觉"), 0)
        with self.assertRaisesRegex(ValueError, "非负整数"):
            parse_score("1.5", "测试/听觉")

    def test_ambiguous_guid_requires_explicit_choice(self) -> None:
        assessment = self.make_assessment(f"{TEST_GUID} {OTHER_GUID}")
        with self.assertRaisesRegex(ValueError, "多个候选"):
            choose_guid(assessment, {})
        self.assertEqual(
            choose_guid(assessment, {assessment.code: TEST_GUID}), TEST_GUID,
        )

    def test_duplicate_guid_requires_content_check(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "first.edf"
            second = Path(directory) / "second.edf"
            first.write_bytes(b"abcd")
            second.write_bytes(b"abce")
            index = {TEST_GUID: (first, second)}
            with self.assertRaisesRegex(ValueError, "内容不同"):
                choose_edf_path(TEST_GUID, index, {})
            second.write_bytes(first.read_bytes())
            self.assertEqual(choose_edf_path(TEST_GUID, index, {}), first)

    def test_healthy_missing_scores_are_expected(self) -> None:
        assessment = Assessment(
            "HC_001_TEST", "TEST", "健康测试",
            dict.fromkeys(SCORE_COLUMNS), TEST_GUID, "测试表", 2,
        )
        self.assertEqual(review_scores(assessment), [])

    def test_workbook_rejects_duplicate_code(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scores.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "编号", "缩写", "患者", *SCORE_COLUMNS, "EDF信息",
            ])
            row = [
                "301-TEST", "TEST", "测试", 0, 0, 2, 0, 0, 0, 2, TEST_GUID,
            ]
            sheet.append(row)
            workbook.save(path)
            assessments = read_assessments(path)
            self.assertEqual(assessments[0].scores["总分"], 2)
            sheet.append(row)
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "编号重复"):
                read_assessments(path)


    def test_cli_keeps_manual_rules_and_patient_ids(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            patient_root = root / "patients"
            healthy_root = root / "healthy"
            patient_root.mkdir()
            healthy_root.mkdir()
            patient_root.joinpath(f"{TEST_GUID}.edf").write_bytes(b"eeg")
            patient_root.joinpath(
                "e4d4b1e4-8e95-43dc-b2a7-a611c01f4dbd.edf"
            ).write_bytes(b"eeg")
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "编号", "缩写", "患者", *SCORE_COLUMNS, "EDF信息",
            ])
            sheet.append([
                "303-MWS", "MWS", "测试", 0, 0, 2, 0, 0, 0, 2,
                f"{TEST_GUID} {OTHER_GUID}",
            ])
            sheet.append([
                "470-MZH", "MZH", "测试", 0, 0, 2, 0, 0, 0, 2, "",
            ])
            sheet.append([
                "359-YSH", "YSH", "测试", 0, 0, 2, 0, 0, 0, 2,
                TEST_GUID,
            ])
            sheet.append([
                "467-YSH", "YSH", "焉石", 3, 2, 3, 3, 1, 3, 15,
                OTHER_GUID,
            ])
            patient_root.joinpath(f"{OTHER_GUID}.edf").write_bytes(b"eeg")
            workbook_path = root / "scores.xlsx"
            workbook.save(workbook_path)
            subprocess.run([
                sys.executable, "-B", str(SCRIPT_PATH),
                "--workbook", str(workbook_path),
                "--patient-root", str(patient_root),
                "--healthy-root", str(healthy_root),
                "--data-root", str(root / "output"),
            ], check=True, capture_output=True, text=True)
            table_path = next((root / "output/tables").glob("01_主表_*"))
            with table_path.open(encoding="utf-8-sig") as stream:
                reader = csv.DictReader(stream)
                rows = list(reader)
            self.assertEqual(len(reader.fieldnames), 14)
            self.assertEqual([row["编号"] for row in rows], [
                "303-MWS", "470-MZH", "467-YSH",
            ])


def identity_row(code, name, total="0"):
    return {"编号": code, "患者": name, "缩写": code.split("-")[-1],
            "总分": total, "GUID": f"guid-{code}", "所在服务器位置": ""}


class PatientIdentityTests(unittest.TestCase):
    def test_repeated_assessments_keep_scores_and_files(self):
        rows = [identity_row("306-LCS", "李春胜", "14"),
                identity_row("447-LCS", "李春胜", "16")]
        result = master_module.assign_patient_ids(rows)
        self.assertEqual([row["患者ID"] for row in result], ["doc_306"] * 2)
        for before, after in zip(rows, result):
            self.assertEqual(before, {key: after[key] for key in before})
            self.assertNotIn("患者ID", before)

    def test_recovery_record_shares_person_identity(self):
        rows = [identity_row("424-WLI", "王立", "14"),
                identity_row("HC_011_WAN", "王立", "")]
        result = master_module.assign_patient_ids(rows)
        self.assertEqual([row["患者ID"] for row in result], ["doc_424"] * 2)
        self.assertEqual(result[1]["编号"], "HC_011_WAN")
        self.assertEqual(result[1]["总分"], "")
        issues = master_module.patient_identity_issues(result)
        self.assertTrue(any("康复后" in issue.reason for issue in issues))

    def test_name_or_initials_do_not_automatically_merge_people(self):
        rows = [identity_row("301-FBI", "同名测试"),
                identity_row("302-FBI", "同名测试")]
        result = master_module.assign_patient_ids(rows)
        self.assertNotEqual(result[0]["患者ID"], result[1]["患者ID"])
        self.assertTrue(master_module.patient_identity_issues(result))

    def test_id_is_stable_when_order_or_cohort_changes(self):
        rows = [identity_row("306-LCS", "李春胜"),
                identity_row("447-LCS", "李春胜")]
        first = master_module.assign_patient_ids(rows)
        updated = master_module.assign_patient_ids(
            [identity_row("300-NEW", "新增患者"), *reversed(rows)]
        )
        old_ids = {row["编号"]: row["患者ID"] for row in first}
        new_ids = {row["编号"]: row["患者ID"] for row in updated}
        for code, patient_id in old_ids.items():
            self.assertEqual(new_ids[code], patient_id)
        subset = master_module.assign_patient_ids(rows[1:])
        self.assertEqual(subset[0]["患者ID"], "doc_306")

    def test_existing_id_survives_missing_anchor_record(self):
        previous = [
            {**identity_row("306-LCS", "李春胜"), "患者ID": "person_fixed"},
            {**identity_row("447-LCS", "李春胜"), "患者ID": "person_fixed"},
        ]
        updated = master_module.assign_patient_ids(
            [identity_row("447-LCS", "李春胜")], previous_rows=previous
        )
        self.assertEqual(updated[0]["患者ID"], "person_fixed")

    def test_cycle_or_name_collision_is_rejected(self):
        rows = [identity_row("301-FBI", "患者甲"),
                identity_row("461-FBI", "患者乙")]
        with self.assertRaisesRegex(ValueError, "不同姓名"):
            master_module.assign_patient_ids(rows)
        cycle = {"301-FBI": "461-FBI", "461-FBI": "301-FBI"}
        with self.assertRaisesRegex(ValueError, "循环"):
            master_module.assign_patient_ids(rows, links=cycle)
        previous = [{**row, "患者ID": "same_id"} for row in rows]
        with self.assertRaisesRegex(ValueError, "不同姓名"):
            master_module.assign_patient_ids(
                rows, links={}, previous_rows=previous
            )


class DiagnosisTests(unittest.TestCase):
    def test_source_labels_keep_clinical_detail(self):
        for label in ["UWS", "MCS+", "MCS-", "eMCS", "UWS-MCS边界"]:
            self.assertEqual(master_module.parse_diagnosis(label, "诊断"), label)
        self.assertEqual(master_module.parse_diagnosis(None, "诊断"), "")

    def test_edf_suffix_is_not_a_diagnosis(self):
        source = f"昏迷-UWS边界神经外科32_{TEST_GUID}"
        self.assertEqual(
            master_module.parse_diagnosis(source, "470-MZH"),
            "昏迷-UWS边界",
        )
        with self.assertRaisesRegex(ValueError, "无法识别"):
            master_module.parse_diagnosis(f"UWS {TEST_GUID}", "诊断")

    def test_corrected_or_blank_labels_survive_rebuild(self):
        row = {**identity_row("301-FBI", "患者甲"), "诊断": "MCS+"}
        for corrected in ["UWS", ""]:
            previous = [{**row, "诊断": corrected}]
            result = master_module.preserve_diagnoses([row], previous)
            self.assertEqual(result[0]["诊断"], corrected)
            self.assertEqual(row["诊断"], "MCS+")

    def test_initial_import_and_new_records_use_source(self):
        row = {**identity_row("301-FBI", "患者甲"), "诊断": "UWS"}
        previous = [{key: value for key, value in row.items() if key != "诊断"}]
        self.assertEqual(
            master_module.preserve_diagnoses([row], previous)[0], row,
        )
        self.assertEqual(master_module.preserve_diagnoses([row], [])[0], row)

    def test_changed_record_and_duplicate_codes_are_rejected(self):
        row = {**identity_row("301-FBI", "患者甲"), "诊断": "UWS"}
        for key in ["患者", "GUID"]:
            previous = [{**row, key: "changed"}]
            with self.assertRaisesRegex(ValueError, "不能复用诊断"):
                master_module.preserve_diagnoses([row], previous)
        with self.assertRaisesRegex(ValueError, "编号重复"):
            master_module.preserve_diagnoses([row], [row, row])

    def test_repeated_visits_keep_separate_labels(self):
        rows = [
            {**identity_row("306-LCS", "李春胜"), "诊断": "UWS"},
            {**identity_row("447-LCS", "李春胜"), "诊断": "MCS+"},
        ]
        rows = master_module.assign_patient_ids(rows)
        result = master_module.preserve_diagnoses(rows, [])
        self.assertEqual([row["诊断"] for row in result], ["UWS", "MCS+"])
        self.assertEqual(result[0]["患者ID"], result[1]["患者ID"])

    def test_workbook_diagnosis_is_read_by_header(self):
        workbook = Workbook()
        sheet = workbook.active
        values = {
            "编号": "301-FBI", "缩写": "FBI", "患者": "患者甲",
            "诊断": "MCS-", **dict.fromkeys(SCORE_COLUMNS, 0),
        }
        assessment = master_module.read_assessment(sheet, 2, values)
        self.assertEqual(assessment.master_values()["诊断"], "MCS-")


if __name__ == "__main__":
    unittest.main()
