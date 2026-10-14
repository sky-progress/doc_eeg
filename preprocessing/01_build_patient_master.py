"""
读取评分 Excel，把每份评分记录对应到一个具体的 EDF 文件，导出含患者ID与诊断的14列主表 CSV 和一份核对报告。
"""

import argparse
import csv
import hashlib
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Mapping
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

SCORE_COLUMNS = (
    "听觉", "视觉", "运动", "口动", "交流", "觉醒", "总分",
)

IDENTITY_COLUMNS = ("编号", "缩写", "患者")

# 导出字段在这里固定，内部追溯信息不额外加入主表。
MASTER_COLUMNS = ("编号", "患者ID", "缩写", "患者") + SCORE_COLUMNS + (
    "诊断", "GUID", "所在服务器位置",
)

MISSING_VALUES = {"", "—", "-"}

GUID_PATTERN = re.compile(
    r"[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}"
)

# 人工确认规则独立列出，避免混入自动匹配判断。
DEFAULT_GUID_OVERRIDES = {
    "303-MWS": "e2ee5b31-758c-46d4-86d8-0b7f0dcb5cd9",
    "470-MZH": "e4d4b1e4-8e95-43dc-b2a7-a611c01f4dbd",
    "463-WFL": "cfd46c86-009a-4b85-903a-334ed69bc0ef",
    "464-WFL": "22a33540-f0d9-4a2c-8f88-ee36eb24a0c7",
}

DEFAULT_EXCLUDED_CODES = {
    "308-WFL", "313-GSR", "359-YSH", "462-MWS",
}

# 用户确认的同人多次测试关系；编号和GUID仍然各自保留。
DEFAULT_PATIENT_LINKS = {
    "461-FBI": "301-FBI",
    "447-LCS": "306-LCS",
    "464-WFL": "463-WFL",
    "465-HAD": "314-HAD",
    "468-WQM": "365-WQM",
    "469-ZS": "372-CZS",
    "470-MZH": "374-MZH",
    "471-ZQD": "397-ZQD",
    "472-ZHR": "399-ZHR",
    "473-ZHR": "399-ZHR",
    "474-GYC": "422-GYC",
    "475-WXL": "442-WXL",
    "HC_011_WAN": "424-WLI",
}

RECOVERED_RECORD_NOTES = {
    "HC_011_WAN": (
        "用户确认是424-WLI王立康复后的记录；二者共用患者ID。"
        "HC目录表示文件来源，后续不能默认将此记录当作无患病史的健康对照。"
    ),
}


# 一条评分记录
@dataclass(frozen=True)
class Assessment:
    code: str                               # 编号，唯一标识一条评分记录
    initials: str                           # 姓名拼音缩写，便于人工核对
    patient_name: str                       # 患者姓名，便于人工核对
    scores: Mapping[str, int | None]        # 六项评分及总分，缺失值保留为None
    edf_description: str                    # 原表EDF信息列，便于人工核对       
    source_sheet: str                       # 评分表工作表名称，便于人工核对
    source_row: int                         # 评分表行号，便于人工核对
    diagnosis: str = ""

    def master_values(self) -> dict:        # 返回评分字段；患者ID和EDF路径在后续流程补齐。
        return {
            "编号": self.code,
            "缩写": self.initials,
            "患者": self.patient_name,
            **self.scores,
            "诊断": self.diagnosis,
        }


@dataclass(frozen=True)
class ReviewIssue:
    code: str
    reason: str


@dataclass(frozen=True)
class BuildResult:
    table_path: str
    report_path: str
    row_count: int
    matched_count: int
    issues: tuple[ReviewIssue, ...]

# 区分缺失值和真实0分，避免后续校验和导出混淆。
def parse_score(value: object, context: str) -> int | None:
    # 缺失与真实0分含义不同，后续校验和导出必须保留这个区别。
    if value is None or str(value).strip() in MISSING_VALUES:
        return None
    try:
        # 接受Excel中的整数文本和2.0，拒绝截断非整数评分。
        number = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError(f"{context}：评分不是有效数字：{value!r}") from error
    if not number.is_finite() or number < 0 or number != number.to_integral():
        raise ValueError(f"{context}：评分必须是非负整数：{value!r}")
    return int(number)

def parse_diagnosis(value: object, context: str) -> str:
    """保留来源诊断含义，只清除误粘的EDF文件信息。"""
    diagnosis = "" if value is None else str(value).strip()
    if diagnosis in MISSING_VALUES:
        return ""
    # 来源470单元格混入科室和GUID；清除格式污染，不推断临床类别。
    edf_suffix = rf"\s*神经外科\d+_{GUID_PATTERN.pattern}\s*$"
    diagnosis = re.sub(edf_suffix, "", diagnosis).strip()
    if GUID_PATTERN.search(diagnosis):
        raise ValueError(f"{context}：诊断含无法识别的EDF信息：{value!r}")
    return diagnosis


# 评分表可能有多个工作表，按字段识别主表，避免依赖可能已过时的工作表名称。
def select_assessment_sheet(workbook):
    # 按字段识别主表，避免依赖可能已过时的工作表名称。
    required = set(IDENTITY_COLUMNS + SCORE_COLUMNS)
    candidates = []
    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(values_only=True), ())
        headers = {str(value).strip() for value in first_row if value}
        if required <= headers:
            candidates.append(sheet)
    if len(candidates) != 1:
        names = [sheet.title for sheet in candidates]
        raise ValueError(f"评分主表应唯一，实际找到：{names}")
    return candidates[0]

# 单行解析，三个身份列任一为空即报错（错误信息带"表名第N行"上下文）
def read_assessment(sheet, row_number: int, values: dict) -> Assessment:
    context = f"{sheet.title}第{row_number}行"
    identity = {}
    for column in IDENTITY_COLUMNS:
        identity[column] = str(values.get(column) or "").strip()
        if not identity[column]:
            raise ValueError(f"{context}：缺少{column}")
    scores = {
        column: parse_score(values.get(column), f"{context}/{column}")
        for column in SCORE_COLUMNS
    }
    description = str(values.get("EDF信息") or "")
    direct_guid = str(values.get("GUID") or "")
    return Assessment(
        code=identity["编号"],
        initials=identity["缩写"],
        patient_name=identity["患者"],
        scores=scores,
        edf_description=f"{direct_guid} {description}".strip(),
        source_sheet=sheet.title,
        source_row=row_number,
        diagnosis=parse_diagnosis(values.get("诊断"), f"{context}/诊断"),
    )

# 读取评分表，返回Assessment列表；编号重复或无数据即报错。
def read_assessments(path: Path) -> list[Assessment]:
    if not path.is_file():
        raise FileNotFoundError(f"评分表不存在：{path}")
    # 读取现有值，不修改原表；公式无缓存时按缺失值处理。
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = select_assessment_sheet(workbook)
        iterator = sheet.iter_rows(values_only=True)
        headers = [
            str(value).strip() if value else ""
            for value in next(iterator)
        ]
        if len([name for name in headers if name]) != len(set(headers) - {""}):
            raise ValueError(f"{sheet.title}：存在重复表头")
        assessments = []
        seen_codes = set()
        for row_number, row in enumerate(iterator, start=2):
            values = dict(zip(headers, row))
            if not values.get("编号"):
                continue
            assessment = read_assessment(sheet, row_number, values)
            if assessment.code in seen_codes:
                raise ValueError(f"{sheet.title}第{row_number}行：编号重复")
            seen_codes.add(assessment.code)
            assessments.append(assessment)
        if not assessments:
            raise ValueError(f"评分主表没有数据：{path}")
        return assessments
    finally:
        workbook.close()

# =================== GUID匹配与EDF路径选择 ===================
def extract_guids(text: str) -> tuple[str, ...]:
    candidates = {
        value.lower() for value in GUID_PATTERN.findall(text)
    }
    return tuple(sorted(candidates))


def build_edf_index(root: Path) -> dict[str, tuple[Path, ...]]:
    if not root.is_dir():
        raise NotADirectoryError(f"原始数据目录不存在：{root}")
    candidates = defaultdict(list)
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() != ".edf":
            continue
        guids = extract_guids(path.name)
        if len(guids) == 1:
            candidates[guids[0]].append(path)
    return {guid: tuple(paths) for guid, paths in candidates.items()}


def choose_guid(assessment: Assessment, overrides: dict[str, str]) -> str:
    candidates = extract_guids(assessment.edf_description)
    # 人工确认优先于旧备注中的候选GUID；没有确认时不猜测。
    if assessment.code in overrides:
        selected = overrides[assessment.code].lower()
        if extract_guids(selected) != (selected,):
            raise ValueError(f"{assessment.code}：人工指定GUID格式错误")
        return selected
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise ValueError("评分表未给出GUID，不按姓名猜测录音")
    raise ValueError(f"评分表有多个候选GUID：{', '.join(candidates)}")


def file_digest(path: Path) -> str:
    hasher = hashlib.sha256()
    try:
        # EDF可能较大，分块计算全文摘要，避免一次载入内存。
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(8 * 1024 * 1024), b""):
                hasher.update(chunk)
    except OSError as error:
        raise OSError(f"无法核验EDF内容：{path}；{error}") from error
    return hasher.hexdigest()


def choose_edf_path(
    guid: str,
    index: dict[str, tuple[Path, ...]],
    digest_cache: dict[Path, str],
) -> Path:
    candidates = index.get(guid, ())
    if not candidates:
        raise FileNotFoundError(f"原始目录中找不到GUID对应EDF：{guid}")
    if len(candidates) == 1:
        return candidates[0]
    # 同GUID和同大小都不能证明内容相同，仍需核对全文摘要。
    if len({path.stat().st_size for path in candidates}) != 1:
        raise ValueError(f"同GUID文件大小不同，不能自动选择：{guid}")
    digests = set()
    for path in candidates:
        if path not in digest_cache:
            digest_cache[path] = file_digest(path)
        digests.add(digest_cache[path])
    if len(digests) != 1:
        raise ValueError(f"同GUID文件内容不同，不能自动选择：{guid}")
    # 内容相同才固定选择路径；不删除或移动其他副本。
    return min(candidates, key=lambda path: str(path))


def review_scores(assessment: Assessment) -> list[ReviewIssue]:
    values = list(assessment.scores.values())
    if assessment.code.lower().startswith("hc_") and all(
        value is None for value in values
    ):
        return []
    if any(value is None for value in values):
        return [ReviewIssue(assessment.code, "评分不完整，缺失值保留为空")]
    # 只报告总分矛盾，不替临床来源表重算或覆盖分数。
    if sum(values[:-1]) != values[-1]:
        return [ReviewIssue(assessment.code, "六项分数之和与原表总分不符")]
    return []


def match_assessment(
    assessment: Assessment,
    index: dict,
    overrides: dict[str, str],
    digest_cache: dict,
) -> tuple[dict, list[ReviewIssue]]:
    row = assessment.master_values()
    row.update({"GUID": "", "所在服务器位置": ""})
    issues = review_scores(assessment)
    try:
        row["GUID"] = choose_guid(assessment, overrides)
        row["所在服务器位置"] = str(
            choose_edf_path(row["GUID"], index, digest_cache)
        )
    except (ValueError, OSError) as error:
        # 单条匹配问题保留在报告中，不让该记录从主表消失。
        issues.append(ReviewIssue(assessment.code, str(error)))
    return row, issues


def canonical_record_code(code: str, links: Mapping[str, str]) -> str:
    """沿明确关联找到患者锚点，拒绝循环关联。"""
    visited = set()
    while code in links:
        if code in visited:
            raise ValueError(f"患者关联存在循环：{code}")
        visited.add(code)
        code = links[code]
    return code


def default_patient_id(code: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", code):
        raise ValueError(f"患者锚点编号格式错误：{code}")
    if code.lower().startswith("hc_"):
        return code.lower()
    number, separator, initials = code.partition("-")
    if not separator or not number.isdigit() or not initials:
        raise ValueError(f"患者记录编号应为数字-缩写：{code}")
    return f"doc_{number}"


def read_existing_master(path: Path | None) -> list[dict]:
    if path is None or not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    codes = [row["编号"] for row in rows]
    if len(set(codes)) != len(codes):
        raise ValueError(f"已有主表编号重复：{path}")
    return rows


def existing_patient_ids(
    previous_rows: list[dict], links: Mapping[str, str],
) -> dict[str, str]:
    candidates = defaultdict(set)
    preferred = {}
    for row in previous_rows:
        patient_id = row.get("患者ID", "")
        if not patient_id:
            continue
        if not re.fullmatch(r"[A-Za-z0-9_-]+", patient_id):
            raise ValueError(f"已有患者ID格式错误：{row['编号']}")
        anchor = canonical_record_code(row["编号"], links)
        candidates[anchor].add(patient_id)
        if anchor == row["编号"]:
            preferred[anchor] = patient_id
    identifiers = {}
    for anchor, values in candidates.items():
        if anchor in preferred:
            identifiers[anchor] = preferred[anchor]
        elif len(values) == 1:
            identifiers[anchor] = next(iter(values))
        else:
            raise ValueError(f"同一患者有多个旧ID且没有锚点记录：{anchor}")
    return identifiers


def assign_patient_ids(
    rows: list[dict], links: Mapping[str, str] | None = None,
    previous_rows: list[dict] | None = None,
) -> list[dict]:
    """只补患者分组键，不改变评分、录音匹配或记录顺序。"""
    links = DEFAULT_PATIENT_LINKS if links is None else links
    previous_rows = previous_rows or []
    identifiers = existing_patient_ids(previous_rows, links)
    previous_names = {row["编号"]: row["患者"] for row in previous_rows}
    names_by_patient_id = defaultdict(set)
    codes = [row["编号"] for row in rows]
    if len(set(codes)) != len(codes):
        raise ValueError("当前主表编号重复，不能分配患者ID")
    assigned = []
    for row in rows:
        code = row["编号"]
        if code in previous_names and previous_names[code] != row["患者"]:
            raise ValueError(f"{code}的姓名与已有主表不同，不能复用身份")
        anchor = canonical_record_code(code, links)
        patient_id = identifiers.get(anchor, default_patient_id(anchor))
        names_by_patient_id[patient_id].add(row["患者"])
        assigned.append({**row, "患者ID": patient_id})
    for patient_id, names in names_by_patient_id.items():
        if len(names) > 1:
            raise ValueError(
                f"同一患者ID对应不同姓名：{patient_id} {sorted(names)}"
            )
    return assigned


def preserve_diagnoses(
    rows: list[dict], previous_rows: list[dict],
) -> list[dict]:
    """同一测试沿用主表现有标签，包含人工修正及主动留空。"""
    previous_by_code = {row["编号"]: row for row in previous_rows}
    if len(previous_by_code) != len(previous_rows):
        raise ValueError("已有主表编号重复，不能复用诊断")
    updated = []
    for row in rows:
        previous = previous_by_code.get(row["编号"])
        diagnosis = row.get("诊断", "")
        if previous is not None and "诊断" in previous:
            # 记录编号若被重新用于另一录音，旧标签不能自动迁移。
            if any(previous[key] != row[key] for key in ("患者", "GUID")):
                raise ValueError(f"{row['编号']}姓名或GUID变化，不能复用诊断")
            diagnosis = previous["诊断"]
        updated.append({**row, "诊断": diagnosis})
    return updated


def patient_identity_issues(rows: list[dict]) -> list[ReviewIssue]:
    identifiers_by_name = defaultdict(set)
    codes_by_name = defaultdict(list)
    issues = []
    for row in rows:
        identifiers_by_name[row["患者"]].add(row["患者ID"])
        codes_by_name[row["患者"]].append(row["编号"])
        if row["编号"] in RECOVERED_RECORD_NOTES:
            issues.append(ReviewIssue(
                row["编号"], RECOVERED_RECORD_NOTES[row["编号"]],
            ))
    # 姓名只是核对线索；新出现的同名记录不自动合并身份。
    for name, identifiers in identifiers_by_name.items():
        if len(identifiers) > 1:
            issues.append(ReviewIssue(
                ", ".join(codes_by_name[name]),
                f"同名记录尚未确认身份关系：{name}，需在划分前核对",
            ))
    return issues


def patient_identity_report_lines(rows: list[dict]) -> list[str]:
    records_by_patient = defaultdict(list)
    for row in rows:
        records_by_patient[row["患者ID"]].append(row)
    lines = [
        "", "## 患者身份与多次测试", "",
        f"患者分组键共{len(records_by_patient)}个；一行仍对应一次测试。",
        "患者ID以稳定记录编号为锚点，不随主表排序或新增记录重新编号。",
        "同一患者所有记录和窗口必须在训练、验证、测试划分中留在同一集合。",
        "康复前后评分各自保留；不同阶段的标签不能未经核对直接平均。", "",
        "| 患者ID | 姓名 | 测试记录 |", "| --- | --- | --- |",
    ]
    for patient_id, records in sorted(records_by_patient.items()):
        if len(records) > 1:
            codes = ", ".join(row["编号"] for row in records)
            lines.append(f"| {patient_id} | {records[0]['患者']} | {codes} |")
    return lines


def write_master(rows: list[dict], path: Path) -> None:
    # 新文件拒绝覆盖；UTF-8 BOM便于Excel正确读取中文。
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=MASTER_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def verify_master(rows: list[dict], path: Path) -> None:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        saved_rows = list(reader)
        if reader.fieldnames != list(MASTER_COLUMNS):
            raise ValueError("导出主表字段与指定的14列不一致")
    expected = [
        {
            key: "" if row[key] is None else str(row[key])
            for key in MASTER_COLUMNS
        }
        for row in rows
    ]
    if saved_rows != expected:
        raise ValueError("主表保存后回读与内存记录不一致")
    for row in saved_rows:
        if row["所在服务器位置"] and not Path(row["所在服务器位置"]).is_file():
            raise FileNotFoundError(f"导出路径不存在：{row['编号']}")


def reused_guid_issues(rows: list[dict]) -> list[ReviewIssue]:
    # 编号对应评分记录，GUID对应录音；重复GUID不直接合并评分。
    codes_by_guid = defaultdict(list)
    for row in rows:
        if row["GUID"]:
            codes_by_guid[row["GUID"]].append(row["编号"])
    return [
        ReviewIssue(
            ", ".join(codes),
            f"多条评分使用同一GUID，保留来源行、不自动合并：{guid}",
        )
        for guid, codes in codes_by_guid.items()
        if len(codes) > 1
    ]


def write_report(
    rows: list[dict], issues: list[ReviewIssue], source: Path, path: Path,
    timestamp: str, roots: tuple[Path, Path],
    excluded_codes: set[str] | None = None,
) -> None:
    matched = sum(bool(row["所在服务器位置"]) for row in rows)
    lines = [
        "# 主表重建报告", "", f"生成时间：{timestamp}",
        f"评分来源：{source}",
        f"评分文件SHA256：{hashlib.sha256(source.read_bytes()).hexdigest()}",
        f"患者EDF目录：{roots[0]}", f"HC EDF目录：{roots[1]}", "",
        f"主表记录：{len(rows)}；找到EDF路径：{matched}；未匹配：{len(rows)-matched}。",
        "评分及编号沿用新版原表；缺失值留空；同GUID副本内容一致后选择固定路径。",
        "未确认排除的记录保留；匹配成功不代表通过信号质控或满足建模纳排条件。",
        "诊断首次取自评分表，仅去除误粘的科室/GUID；不按总分推算诊断。",
        "已有诊断按测试编号保留，包含人工修正或主动留空；新增记录取来源标签。",
        f"用户确认不纳入：{', '.join(sorted(excluded_codes or [])) or '无'}",
        "", "## 需要核对的记录", "",
    ]
    if issues:
        lines.extend(f"- {issue.code}：{issue.reason}" for issue in issues)
    else:
        lines.append("无。")
    lines.extend(patient_identity_report_lines(rows))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def match_assessments(
    assessments: list[Assessment], patient_root: Path, healthy_root: Path,
    overrides: Mapping[str, str],
) -> tuple[list[dict], list[ReviewIssue]]:
    # 两组分开索引，防止在错误的数据区域找到同名文件。
    patient_index = build_edf_index(patient_root)
    healthy_index = build_edf_index(healthy_root)
    rows, issues, digest_cache = [], [], {}
    for assessment in assessments:
        is_healthy = assessment.code.lower().startswith("hc_")
        edf_index = healthy_index if is_healthy else patient_index
        row, record_issues = match_assessment(
            assessment, edf_index, overrides or {}, digest_cache,
        )
        rows.append(row)
        issues.extend(record_issues)
        print(f"主表匹配：{assessment.code}", flush=True)
    return rows, issues


def build_master(
    workbook_path: Path, patient_root: Path, healthy_root: Path,
    table_directory: Path, report_directory: Path,
    overrides: dict[str, str] | None = None,
    excluded_codes: set[str] | None = None,
    patient_links: Mapping[str, str] | None = None,
    previous_master: Path | None = None,
) -> BuildResult:
    # 仅执行明确确认的排除，不沿用旧实验的纳排名单。
    assessments = [
        assessment for assessment in read_assessments(workbook_path)
        if assessment.code not in (excluded_codes or set())
    ]
    rows, issues = match_assessments(
        assessments, patient_root, healthy_root, overrides or {}
    )
    previous_rows = read_existing_master(previous_master)
    rows = assign_patient_ids(rows, patient_links, previous_rows)
    rows = preserve_diagnoses(rows, previous_rows)
    issues.extend(reused_guid_issues(rows))
    issues.extend(patient_identity_issues(rows))
    now = datetime.now(ZoneInfo("Asia/Shanghai"))
    timestamp = now.strftime("%Y%m%d_%H%M%S")
    table_directory.mkdir(parents=True, exist_ok=True)
    report_directory.mkdir(parents=True, exist_ok=True)
    table_path = table_directory / f"01_主表_{timestamp}.csv"
    report_path = report_directory / f"01_主表重建_{timestamp}.md"
    # 回读检查实际保存内容，不能仅凭写文件成功判定主表正确。
    write_master(rows, table_path)
    verify_master(rows, table_path)
    write_report(
        rows, issues, workbook_path, report_path,
        now.isoformat(timespec="seconds"), (patient_root, healthy_root),
        excluded_codes,
    )
    return BuildResult(
        str(table_path), str(report_path), len(rows),
        sum(bool(row["所在服务器位置"]) for row in rows), tuple(issues),
    )


def parse_guid_override(value: str) -> tuple[str, str]:
    code, separator, guid = value.partition("=")
    if not separator or not code.strip() or not guid.strip():
        raise argparse.ArgumentTypeError("GUID覆盖格式应为 编号=GUID")
    return code.strip(), guid.strip().lower()


def parse_patient_link(value: str) -> tuple[str, str]:
    code, separator, anchor = value.partition("=")
    if not separator or not code.strip() or not anchor.strip():
        raise argparse.ArgumentTypeError("患者关联格式应为 编号=锚点编号")
    return code.strip(), anchor.strip()


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="构建精简患者主表")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--patient-root", type=Path, required=True)
    parser.add_argument("--healthy-root", type=Path, required=True)
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument(
        "--guid-override", action="append", type=parse_guid_override,
        default=[], help="仅用于用户明确确认的评分与GUID对应",
    )
    parser.add_argument(
        "--exclude-code", action="append", default=[],
        help="用户确认不纳入主表的编号，不删除原始数据",
    )
    parser.add_argument(
        "--patient-link", action="append", type=parse_patient_link,
        default=[], help="明确同一人的记录关联：测试编号=患者锚点编号",
    )
    parser.add_argument(
        "--previous-master", type=Path,
        help="沿用已有患者ID及诊断；诊断列存在时保留人工修改与空值",
    )
    return parser.parse_args()


def main() -> None:
    arguments = parse_arguments()
    supplied_overrides = dict(arguments.guid_override)
    if len(supplied_overrides) != len(arguments.guid_override):
        raise ValueError("同一编号重复指定GUID，需明确唯一对应")
    overrides = {**DEFAULT_GUID_OVERRIDES, **supplied_overrides}
    excluded_codes = DEFAULT_EXCLUDED_CODES | set(arguments.exclude_code)
    supplied_links = dict(arguments.patient_link)
    if len(supplied_links) != len(arguments.patient_link):
        raise ValueError("同一记录重复指定患者关联")
    patient_links = {**DEFAULT_PATIENT_LINKS, **supplied_links}
    previous_master = arguments.previous_master or (
        arguments.data_root / "tables/63_主表_20261013_165853.csv"
    )
    result = build_master(
        arguments.workbook, arguments.patient_root,
        arguments.healthy_root, arguments.data_root / "tables",
        arguments.data_root / "reports", overrides, excluded_codes,
        patient_links, previous_master,
    )
    print(f"主表：{result.table_path}")
    print(f"报告：{result.report_path}")
    print(f"记录{result.row_count}条，匹配{result.matched_count}条")


if __name__ == "__main__":
    main()
