"""只读检查EDF头部、通道字段和数据区，不修复或转换原始文件。"""
import argparse,csv,hashlib,html,json,math,warnings
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
import mne
from openpyxl import load_workbook
from reporting import write_edf_report

FIELDS=[('label',16),('transducer',80),('unit',8),('physical_min',8),('physical_max',8),('digital_min',8),('digital_max',8),('prefilter',80),('samples_per_record',8),('reserved',32)]
NUMBERS={'physical_min':float,'physical_max':float,'digital_min':float,'digital_max':float,'samples_per_record':int}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fixed_header(head,shift):
    """偏移是只读解析假设，不能作为自动修复结论。"""
    def value(start,end):return head[start+shift:end+shift].decode('ascii').strip()
    try:
        r={'shift':shift,'header_bytes':int(value(184,192)),'records':int(value(236,244)),'seconds':float(value(244,252)),'channels':int(value(252,256)),'date_raw':value(168,176),'time_raw':value(176,184)}
        return r if 0<r['channels']<1024 and r['header_bytes']>=256 and r['records']>=-1 and math.isfinite(r['seconds']) and r['seconds']>0 else None
    except (ValueError,UnicodeDecodeError):return None


def inspect_file(path,code,group):
    before=path.stat()
    r={'code':code,'group':group,'path':str(path),'file_bytes':before.st_size,'mtime_ns':before.st_mtime_ns,'layout_shift':None,'n_channels':None,'header_bytes_declared':None,'header_candidate':None,'records_declared':None,'record_seconds':None,'records_by_size':None,'bytes_per_record':None,'remainder_bytes':None,'duration_seconds':None,'channel_parse_errors':0,'channel_warnings':[],'issues':[],'mne_readable':False,'mne_error':None,'mne_warnings':[]}
    channels=[]
    try:
        with path.open('rb') as f:
            head=f.read(512)
            r['version_raw']=head[:8].decode('latin-1').strip()
            candidates=[v for shift in [0,-80,80] if (v:=fixed_header(head,shift)) is not None]
            r['candidate_layouts']=candidates
            if len(candidates)!=1:r['issues'].append('固定头候选布局缺失或不唯一')
            else:
                fixed=candidates[0];n=fixed['channels'];shift=fixed['shift'];start=256+shift;actual=start+256*n
                r.update(layout_shift=shift,n_channels=n,header_bytes_declared=fixed['header_bytes'],header_candidate=actual,records_declared=fixed['records'],record_seconds=fixed['seconds'],date_raw=fixed['date_raw'],time_raw=fixed['time_raw'])
                if fixed['header_bytes']!=256*(n+1):r['issues'].append('声明头长度不符合标准通道数计算')
                if shift:r['issues'].append(f'头部关键字段和通道起点相对标准偏移{shift}字节')
                f.seek(start);block=f.read(256*n)
                if len(block)!=256*n:r['issues'].append('通道头不完整')
                else:
                    channels=[{'code':code,'group':group,'channel_index':i+1} for i in range(n)];offset=0
                    for name,width in FIELDS:
                        for i in range(n):
                            text=block[offset+i*width:offset+(i+1)*width].decode('latin-1').strip()
                            channels[i][name+'_raw']=text
                            if name in NUMBERS:
                                try:
                                    number=NUMBERS[name](text)
                                    if not math.isfinite(number):raise ValueError('非有限数')
                                    channels[i][name]=number
                                except ValueError:
                                    channels[i][name]=None;r['channel_parse_errors']+=1
                                    r['issues'].append(f'通道{i+1}的{name}不是有效数字')
                            else:channels[i][name]=text
                        offset+=n*width
                    labels=[c['label'] for c in channels]
                    if any(not x or any(ord(ch)<32 for ch in x) for x in labels):r['channel_warnings'].append('空白或控制字符通道名')
                    if len(set(labels))!=len(labels):r['channel_warnings'].append('通道名重复')
                    for c in channels:
                        # 注释通道的标定范围不按生理信号标准判断。
                        if 'annotations' not in c['label'].lower():
                            # EDF允许负增益：物理max小于min不表示损坏，不能擅自交换。
                            if c.get('physical_min') is not None and c.get('physical_max') is not None:
                                if c['physical_max']==c['physical_min']:r['channel_warnings'].append(f"通道{c['channel_index']}的物理范围为零")
                                c['negative_gain']=c['physical_max']<c['physical_min']
                            if c.get('digital_min') is not None and c.get('digital_max') is not None and c['digital_max']<=c['digital_min']:
                                r['channel_warnings'].append(f"通道{c['channel_index']}的数字范围无效")
                    samples=[c.get('samples_per_record') for c in channels]
                    if all(v is not None and v>0 for v in samples):
                        size=2*sum(samples);data_bytes=before.st_size-actual
                        count,remainder=divmod(data_bytes,size) if data_bytes>=0 else (0,data_bytes)
                        r.update(bytes_per_record=size,records_by_size=count,remainder_bytes=remainder,duration_seconds=count*fixed['seconds'])
                        # 标准记录数位置若与实际记录数相符且落入标签区，记录覆盖证据。
                        text=head[236:244].decode('latin-1')
                        r['record_count_standard_offset_raw']=text
                        try:r['record_count_overlaps_channel_labels']=shift==-80 and int(text.strip())==count
                        except ValueError:r['record_count_overlaps_channel_labels']=False
                        if r['record_count_overlaps_channel_labels']:
                            r['channel_warnings'].append('记录数被写入标准偏移，覆盖第4、5通道标签')
                        if remainder:r['issues'].append('候选数据区存在不完整记录或额外尾部字节')
                        if fixed['records']!=-1 and fixed['records']!=count:r['issues'].append('声明记录数与文件大小推算不一致')
                        # 仅抽查首末记录长度，不加载完整波形，不评估伪迹。
                        if count>0:
                            f.seek(actual);first=len(f.read(size));f.seek(actual+(count-1)*size);last=len(f.read(size))
                            r['first_last_bytes_read']=[first,last]
                            if first!=size or last!=size:r['issues'].append('首末数据记录读取不完整')
                        else:r['issues'].append('没有完整记录')
                        for c,value in zip(channels,samples):c['sampling_rate_hz']=value/fixed['seconds']
                        r['sampling_rates_hz']=sorted(set(c['sampling_rate_hz'] for c in channels))
                    else:r['issues'].append('每记录样本数缺失或非正数，无法可靠确定数据区')
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            try:
                raw=mne.io.read_raw_edf(str(path),preload=False,verbose='ERROR')
                r.update(mne_readable=True,mne_sfreq_hz=raw.info['sfreq'],mne_n_channels=len(raw.ch_names),mne_n_samples=raw.n_times);raw.close()
            except Exception as error:r['mne_error']=type(error).__name__+': '+str(error)
            r['mne_warnings']=[str(w.message) for w in caught]
        if r['layout_shift'] is None:r['status']='无法确定头部布局'
        elif r['channel_parse_errors'] or r['bytes_per_record'] is None:r['status']='通道数字字段异常，需单独复核'
        elif r['remainder_bytes']!=0 or not r['records_by_size']:r['status']='数据区长度异常，需单独复核'
        elif r['header_bytes_declared']!=256*(r['n_channels']+1):r['status']='声明头长度异常，需单独复核'
        elif r.get('record_count_overlaps_channel_labels'):r['status']='固定头少80字节，记录数覆盖通道标签'
        elif r['channel_warnings']:r['status']='通道范围或标签异常，需单独复核'
        elif r['layout_shift']==-80:r['status']='非标准布局，结构支持固定头少80字节'
        elif r['layout_shift']:r['status']='其他非标准布局，需单独复核'
        else:r['status']='标准布局，仍需信号质控'
    except Exception as error:r.update(status='检查异常，需单独复核');r['issues'].append(type(error).__name__+': '+str(error))
    after=path.stat()
    assert (before.st_size,before.st_mtime_ns)==(after.st_size,after.st_mtime_ns),'检查过程中源文件发生变化'
    return r,channels


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--master',type=Path,required=True);parser.add_argument('--logs',type=Path,required=True);parser.add_argument('--tables',type=Path,required=True);parser.add_argument('--reports',type=Path,required=True);parser.add_argument('--codes',nargs='+',help='仅检查指定评分编号');args=parser.parse_args()
    book=load_workbook(args.master,read_only=True,data_only=True);rows=list(book['受试者主表'].iter_rows(values_only=True));book.close();cols={v:i for i,v in enumerate(rows[0])}
    if args.codes:
        selected=set(args.codes);rows=[rows[0]]+[r for r in rows[1:] if str(r[cols['编号']]) in selected]
        assert {str(r[cols['编号']]) for r in rows[1:]}==selected,'指定编号未在主表中全部找到'
    assert len({r[cols['服务器EDF路径']] for r in rows[1:]})==len(rows)-1,'主表路径重复'
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d_%H%M%S')
    run=args.logs/('02_inspect_'+stamp);run.mkdir(parents=True,exist_ok=False)
    args.tables.mkdir(parents=True,exist_ok=True)
    details=[];channels=[]
    with (run/'log.txt').open('x',encoding='utf-8') as log:
        for i,row in enumerate(rows[1:],1):
            result,channel=inspect_file(Path(row[cols['服务器EDF路径']]),str(row[cols['编号']]),str(row[cols['组别']]))
            details.append(result);channels.extend(channel)
            line=json.dumps({'index':i,'total':len(rows)-1,'code':result['code'],'status':result['status']},ensure_ascii=False);log.write(line+'\n');log.flush()
            if i%20==0 or i==len(rows)-1:print(line,flush=True)
    with (run/'details.jsonl').open('x',encoding='utf-8') as f:
        for item in details:f.write(json.dumps(item,ensure_ascii=False)+'\n')
    file_cols=['code','group','path','file_bytes','layout_shift','n_channels','header_bytes_declared','header_candidate','records_declared','record_seconds','records_by_size','bytes_per_record','remainder_bytes','duration_seconds','channel_parse_errors','mne_readable','mne_error','status']
    for name,records,fieldnames in [('files.csv',details,file_cols),('channels.csv',channels,list(dict.fromkeys(k for c in channels for k in c)))]:
        with (args.tables/(('02_文件检查_' if name=='files.csv' else '02_通道字段_')+stamp+'.csv')).open('x',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=fieldnames,extrasaction='ignore');writer.writeheader();writer.writerows(records)
    counts=Counter(r['status'] for r in details)
    summary={'files_checked':len(details),'group_counts':dict(Counter(r['group'] for r in details)),'status_counts':dict(counts),'mne_readable_count':sum(r['mne_readable'] for r in details),'channel_rows':len(channels),'channel_count_distribution':dict(Counter(str(r['n_channels']) for r in details)),'layout_shift_distribution':dict(Counter(str(r['layout_shift']) for r in details)),'data_region_aligned_count':sum(r['remainder_bytes']==0 for r in details),'record_count_overlaps_labels_count':sum(r.get('record_count_overlaps_channel_labels',False) for r in details),'channel_numeric_error_count':sum(r['channel_parse_errors'] for r in details),'master':str(args.master),'master_sha256':digest(args.master),'script':str(Path(__file__)),'script_sha256':digest(Path(__file__)),'mne_version':mne.__version__,'run_directory':str(run),'source_files_modified':False,'repair_performed':False,'signal_quality_assessed':False,'created_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat()}
    assert len(details)==len(rows)-1
    (run/'config.json').write_text(json.dumps({'master':str(args.master),'candidate_fixed_header_shifts':[0,-80,80],'preload':False,'repair':False},ensure_ascii=False,indent=2),encoding='utf-8')
    (run/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    summary['report']=write_edf_report(details,summary,args.reports,stamp)
    (run/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print('FINAL_SUMMARY '+json.dumps(summary,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
