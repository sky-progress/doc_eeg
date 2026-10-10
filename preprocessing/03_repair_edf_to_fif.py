"""最小修复EDF结构并转FIF；保留源数据和负增益，EEG使用10–20坐标。"""
import argparse,csv,hashlib,html,json,re,warnings
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
import mne
import numpy as np
from openpyxl import load_workbook
from reporting import write_report
from mne.io.edf.edf import _read_annotations_edf

EEG19=['Fp1','Fp2','F3','F4','C3','C4','P3','P4','O1','O2','F7','F8','T3','T4','T5','T6','Fz','Cz','Pz']
EEG={'Fp1','Fp2','F3','F4','C3','C4','P3','P4','O1','O2','F7','F8','T3','T4','T5','T6','A1','A2','Fz','Cz','Pz'}
TAL=re.compile(rb'^([+-]\d+(?:\.\d*)?)\x14\x14\x00')
COLUMNS=['code','group','guid','source_edf','source_bytes','source_mtime_ns','repaired_edf','fif_files','status','error','n_channels','n_samples','sfreq_hz','annotations','annotation_onset_max_error_s','annotation_duration_max_error_s','data_sha256','timekeeping_continuous','time_gaps','fif_timeline','eeg19_status','channel_names','boundary_count','original_annotation_count','validation','fif_bytes']


def checkpoint(path,records):
    """同一对应表逐步更新，不生成多个重复版本。"""
    temporary=path.with_suffix('.csv.tmp')
    with temporary.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=COLUMNS,extrasaction='ignore');writer.writeheader();writer.writerows(records.values())
    temporary.replace(path)


def repair(source,target,inspection):
    """只补缺失recording字段并修复覆盖标签，数据区逐字节复制。"""
    n=inspection['n_channels'];actual=176+256*n;standard=256+256*n;record_size=inspection['bytes_per_record'];count=inspection['records_by_size']
    with source.open('rb') as f:old=f.read(actual)
    assert inspection['layout_shift']==-80 and inspection['remainder_bytes']==0
    tokens=old[8:88].decode('latin-1').split()
    assert len(tokens)>=4 and tokens[1] in ['M','F','X'] and re.fullmatch(r'(\d{2}-[A-Za-z]{3}-\d{4}|X)',tokens[2]),'剩余身份字段不符合patient identification，停止猜测'
    labels=[old[176+i*16:176+(i+1)*16].decode('latin-1').strip() for i in range(n)]
    assert labels[:3]==['Fp1','Fp2','F3'] and labels[5]=='C4'
    assert labels[3].startswith('F4') and all(ch.isdigit() or ch.isspace() for ch in labels[4]),'被覆盖通道不符合已确认模板'
    date=old[88:96].decode('ascii');day,month,year=map(int,date.split('.'));year+=2000 if year<85 else 1900
    month_name=['JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP','OCT','NOV','DEC'][month-1]
    # 缺失的采集者、设备等信息用X表示未知，不能虚构。
    recording=f'Startdate {day:02d}-{month_name}-{year:04d} X X X'.encode().ljust(80,b' ')
    header=bytearray(old[:88]+recording+old[88:]);assert len(header)==standard
    header[236:244]=str(count).encode().ljust(8,b' ')
    header[256+3*16:256+4*16]=b'F4'.ljust(16,b' ')
    header[256+4*16:256+5*16]=b'C3'.ljust(16,b' ')
    samples_at=176+n*(16+80+8+8+8+8+8+80)
    samples=[int(old[samples_at+i*8:samples_at+(i+1)*8]) for i in range(n)]
    annotation=[i for i,label in enumerate(labels) if label=='EDF Annotations']
    assert len(annotation)<=1,'多注释通道需单独确认'
    if annotation:header[192:236]=b'EDF+C'.ljust(44,b' ')
    digest=hashlib.sha256();first_onset=None;previous_onset=None;gaps=[];seen=0
    with source.open('rb') as f,target.open('xb') as out:
        f.seek(actual);out.write(header)
        batch_records=max(1,(8*1024*1024)//record_size)
        while block:=f.read(batch_records*record_size):
            assert len(block)%record_size==0
            if annotation:
                i=annotation[0];offset=2*sum(samples[:i]);length=2*samples[i]
                for start in range(0,len(block),record_size):
                    data=block[start+offset:start+offset+length]
                    data.decode('utf-8')  # 无法正确解码时保留原文件并报告，不静默替换事件文字。
                    match=TAL.match(data);assert match,'缺少有效EDF+时间标记'
                    onset=float(match.group(1))
                    if first_onset is None:first_onset=onset
                    if previous_onset is not None:
                        difference=onset-previous_onset-inspection['record_seconds']
                        assert difference>=-1e-6,'记录时间倒退，需单独复核'
                        if difference>1e-6:gaps.append({'record':seen,'gap_seconds':difference,'next_onset_seconds':onset})
                    previous_onset=onset;seen+=1
            digest.update(block);assert out.write(block)==len(block)
    if gaps:
        with target.open('r+b') as f:f.seek(192);f.write(b'EDF+D'.ljust(44,b' '))
    assert target.stat().st_size==source.stat().st_size+80
    # 从输出文件重新计算数据区摘要，验证复制内容真实一致。
    output_hash=hashlib.sha256()
    with target.open('rb') as f:
        f.seek(standard)
        while block:=f.read(8*1024*1024):output_hash.update(block)
    assert output_hash.hexdigest()==digest.hexdigest(),'修复前后数据区不一致'
    assert not annotation or seen==count
    return digest.hexdigest(),gaps


def channel_types(names):
    return {name:('eeg' if name in EEG else 'eog' if name in {'ROC','LOC'} else 'ecg' if name=='EKG' else 'emg' if name.startswith('EMG') else 'misc') for name in names}


def validate_digital(source,raw,inspection):
    """用源文件数字值独立验证MNE标定、极性及通道顺序。"""
    n=inspection['n_channels'];actual=176+256*n
    with source.open('rb') as f:f.seek(176);h=f.read(256*n)
    fields={};offset=0
    for name,width in [('label',16),('transducer',80),('unit',8),('pmin',8),('pmax',8),('dmin',8),('dmax',8),('prefilter',80),('samples',8),('reserved',32)]:
        fields[name]=[h[offset+i*width:offset+(i+1)*width].decode('latin-1').strip() for i in range(n)];offset+=n*width
    samples=list(map(int,fields['samples']));selected=[i for i,label in enumerate(fields['label']) if label!='EDF Annotations']
    assert all(samples[i]==int(raw.info['sfreq']*inspection['record_seconds']) for i in selected),'信号通道混合采样率需单独确认'
    for record in sorted({0,inspection['records_by_size']//2,inspection['records_by_size']-1}):
        with source.open('rb') as f:f.seek(actual+record*inspection['bytes_per_record']);block=f.read(inspection['bytes_per_record'])
        digital=np.frombuffer(block,dtype='<i2');expected=[];start=0
        for i,size in enumerate(samples):
            if i in selected:
                pmin,pmax,dmin,dmax=[float(fields[key][i]) for key in ['pmin','pmax','dmin','dmax']]
                physical=pmin+(digital[start:start+size].astype(float)-dmin)*(pmax-pmin)/(dmax-dmin)
                unit=fields['unit'][i];scale=1e-6 if unit in ['uV','µV','μV'] else 1e-3 if unit=='mV' else 1
                expected.append(physical*scale)
            start+=size
        sample=record*samples[selected[0]]
        np.testing.assert_allclose(raw.get_data(start=sample,stop=sample+samples[selected[0]]),np.array(expected),rtol=1e-12,atol=1e-12)


def joined_annotations(edf,raw,gaps,inspection):
    """从完整TAL恢复事件，把采集暂停压缩为边界，不丢失MNE裁剪前事件。"""
    n=inspection['n_channels']
    with Path(edf).open('rb') as f:
        fixed=f.read(256);header_bytes=int(fixed[184:192]);count=int(fixed[236:244]);h=f.read(256*n)
        labels=[h[i*16:(i+1)*16].decode('latin-1').strip() for i in range(n)]
        i=labels.index('EDF Annotations');at=n*(16+80+8+8+8+8+8+80)
        samples=[int(h[at+j*8:at+(j+1)*8]) for j in range(n)]
        size=2*sum(samples);offset=2*sum(samples[:i]);length=2*samples[i]
        payload=bytearray();first=None;records=0
        while block:=f.read(max(1,(8*1024*1024)//size)*size):
            assert len(block)%size==0
            for start in range(0,len(block),size):
                tal=block[start+offset:start+offset+length]
                if first is None:
                    match=TAL.match(tal);assert match;first=float(match.group(1))
                payload.extend(tal);records+=1
        assert records==count
    annotations=_read_annotations_edf(np.frombuffer(payload,dtype='<i2'),ch_names=raw.ch_names)
    intervals=[(g['next_onset_seconds']-first-g['gap_seconds'],g['next_onset_seconds']-first) for g in gaps]
    def compress(values):
        result=np.asarray(values,dtype=float).copy()
        for start,end in intervals:result-=np.clip(np.asarray(values)-start,0,end-start)
        return result
    for gap,(start,end) in zip(gaps,intervals):
        np.testing.assert_allclose(compress([end])[0],gap['record']*inspection['record_seconds'],rtol=0,atol=1e-6)
    onsets=compress(annotations.onset)
    durations=np.maximum(0,compress(annotations.onset+annotations.duration)-onsets)
    signal_end=raw.n_times/raw.info['sfreq']
    extras=[{'original_onset_s':float(onset),'original_duration_s':float(duration),'in_removed_gap':bool(any(start<=onset<end for start,end in intervals)),'outside_signal_span':bool(new_onset<0 or new_onset>signal_end)} for onset,duration,new_onset in zip(annotations.onset,annotations.duration,onsets)]
    # append不裁剪范围外的原始记录事件；用extras标记，保留完整可追溯信息。
    raw.set_annotations(mne.Annotations([],[],[],orig_time=raw.info['meas_date']))
    raw.annotations.append(onsets,durations,annotations.description,ch_names=annotations.ch_names,extras=extras)
    for gap in gaps:
        onset=gap['record']*inspection['record_seconds']
        extra={'original_gap_seconds':float(gap['gap_seconds']),'original_next_onset_s':float(gap['next_onset_seconds']),'record_index':int(gap['record'])}
        raw.annotations.append([onset,onset],[0,0],['BAD boundary','EDGE boundary'],extras=[extra,extra])
    assert len(raw.annotations)==len(annotations)+2*len(gaps)
    return len(annotations)


def convert(edf,stage,code,source,inspection,joined_gaps=None):
    raw=mne.io.read_raw_edf(edf,preload=False,verbose='ERROR')
    try:
        assert raw.ch_names[3:5]==['F4','C3'];assert raw.n_times==inspection['records_by_size']*int(raw.info['sfreq']*inspection['record_seconds'])
        validate_digital(source,raw,inspection)
        original_annotation_count=joined_annotations(edf,raw,joined_gaps,inspection) if joined_gaps else len(raw.annotations)
        raw.set_channel_types(channel_types(raw.ch_names),verbose='ERROR')
        raw.set_montage(mne.channels.make_standard_montage('standard_1020'),match_case=False,on_missing='raise',verbose='ERROR')
        for channel in raw.info['chs']:
            if channel['ch_name'] in EEG:assert np.isfinite(channel['loc'][:3]).all()
        event_descriptions=raw.annotations.description.copy()
        raw.pick(EEG19)
        np.testing.assert_array_equal(raw.annotations.description,event_descriptions)
        first=stage/(code+'_raw.fif')
        if first.exists():
            try:
                existing=mne.io.read_raw_fif(first,preload=False,verbose='ERROR')
                complete=existing.n_times==raw.n_times;existing.close()
            except Exception:complete=False
            if not complete:
                for temporary in stage.glob(code+'_raw*.fif'):
                    assert temporary.resolve().parent==stage.resolve();temporary.unlink()
        if first.exists():
            files=sorted(stage.glob(code+'_raw*.fif'),key=lambda f:(f!=first,f.name))
        else:
            files=raw.save(first,fmt='single',buffer_size_sec=30,split_size='2GB',overwrite=False,verbose='ERROR')
        check=mne.io.read_raw_fif(files[0],preload=False,verbose='ERROR')
        try:
            assert check.ch_names==raw.ch_names and check.n_times==raw.n_times and check.info['sfreq']==raw.info['sfreq']
            assert check.get_channel_types()==raw.get_channel_types()
            np.testing.assert_array_equal([html.unescape(value) for value in check.annotations.description],raw.annotations.description)
            # FIF保存float32起止时间；核验精确的存储舍入，而不是任意放宽容差。
            expected_onsets=raw.annotations.onset.astype('float32').astype(float)
            expected_ends=(raw.annotations.onset+raw.annotations.duration).astype('float32').astype(float)
            np.testing.assert_allclose(check.annotations.onset,expected_onsets,rtol=0,atol=1e-6)
            np.testing.assert_allclose(check.annotations.duration,expected_ends-expected_onsets,rtol=0,atol=1e-6)
            assert check.info['meas_date']==raw.info['meas_date']
            assert [dict(extra) for extra in check.annotations.extras]==[dict(extra) for extra in raw.annotations.extras]
            if joined_gaps:
                assert sum(value=='BAD boundary' for value in check.annotations.description)==len(joined_gaps)
            for start in sorted({0,max(0,raw.n_times//2-250),max(0,raw.n_times-750)}):
                stop=min(raw.n_times,start+750)
                # FIF单精度是MNE默认存储方式，核对其预期float32舍入。
                np.testing.assert_allclose(check.get_data(start=start,stop=stop),raw.get_data(start=start,stop=stop).astype('float32').astype(float),rtol=0,atol=0)
            result={'n_channels':len(raw.ch_names),'eeg19_status':'完成','channel_names':json.dumps(raw.ch_names),'n_samples':raw.n_times,'sfreq_hz':raw.info['sfreq'],'annotations':len(raw.annotations),'original_annotation_count':original_annotation_count,'boundary_count':len(joined_gaps or []),'fif_timeline':'压缩采集暂停，保留BAD/EDGE边界' if joined_gaps else '原连续时间轴','annotation_onset_max_error_s':float(np.max(np.abs(check.annotations.onset-raw.annotations.onset),initial=0)),'annotation_duration_max_error_s':float(np.max(np.abs(check.annotations.duration-raw.annotations.duration),initial=0)),'validation':'数据区SHA256一致；源数字标定一致；FIF首中末样本、通道、事件及时间一致'}
        finally:check.close()
        return [Path(f) for f in files],result
    finally:raw.close()


def select_eeg19(args):
    """只筛选现有FIF；新文件验证通过后替换，失败可恢复原文件。"""
    base=args.base;folder=(base/'repaired/fif').resolve();manifest=base/'tables/03_修复对应表.csv'
    with manifest.open(encoding='utf-8-sig',newline='') as f:records={r['code']:r for r in csv.DictReader(f)}
    selected=[code for code in records if not args.codes or code in set(args.codes)]
    if args.codes:assert set(selected)==set(args.codes)
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d_%H%M%S')
    staging=base/'repaired/.select19';staging.mkdir(exist_ok=True)
    failures=[]
    with (base/'logs'/('03_19导筛选_'+stamp+'.jsonl')).open('x',encoding='utf-8') as log:
        for index,code in enumerate(selected,1):
            row=records[code]
            if row.get('eeg19_status')=='完成' and int(row['n_channels'])==19:
                assert all(Path(f).is_file() for f in json.loads(row['fif_files']))
                print(json.dumps({'index':index,'code':code,'status':'沿用19导FIF'},ensure_ascii=False),flush=True);continue
            original=row.copy();old=[Path(f).resolve() for f in json.loads(row['fif_files'])]
            assert all(f.parent==folder and f.name.startswith(code+'_raw') and f.suffix=='.fif' for f in old)
            work=staging/code;newdir=work/'new';backup=work/'old'
            assert not work.exists(),'存在未完成筛选，先核对后恢复'
            newdir.mkdir(parents=True);backup.mkdir()
            newfiles=[];published=[];moved=[];raw=None;picked=None;check=None
            try:
                raw=mne.io.read_raw_fif(old[0],preload=False,verbose='ERROR')
                assert set(EEG19)<=set(raw.ch_names),'缺少目标EEG通道'
                picked=raw.copy().pick(EEG19)
                np.testing.assert_array_equal(picked.annotations.description,raw.annotations.description)
                assert list(picked.annotations.ch_names)==list(raw.annotations.ch_names),'删除通道会影响关联事件，停止替换'
                newfiles=[Path(f) for f in picked.save(newdir/(code+'_raw.fif'),fmt='single',buffer_size_sec=30,split_size='2GB',overwrite=False,verbose='ERROR')]
                check=mne.io.read_raw_fif(newfiles[0],preload=False,verbose='ERROR')
                assert check.ch_names==EEG19 and check.n_times==raw.n_times and check.info['sfreq']==raw.info['sfreq']
                assert check.info['meas_date']==raw.info['meas_date']
                np.testing.assert_array_equal(check.annotations.description,raw.annotations.description)
                assert [dict(e) for e in check.annotations.extras]==[dict(e) for e in raw.annotations.extras]
                np.testing.assert_allclose(check.annotations.onset,raw.annotations.onset,rtol=0,atol=1e-6)
                np.testing.assert_allclose(check.annotations.duration,raw.annotations.duration,rtol=0,atol=1e-6)
                points={0,max(0,raw.n_times//2-250),max(0,raw.n_times-750)}
                for onset,description in zip(raw.annotations.onset,raw.annotations.description):
                    if description=='BAD boundary':points.add(max(0,min(raw.n_times-1,int(round(onset*raw.info['sfreq']))-250)))
                for start in sorted(points):
                    stop=min(raw.n_times,start+750)
                    np.testing.assert_array_equal(check.get_data(start=start,stop=stop),raw.get_data(picks=EEG19,start=start,stop=stop))
                np.testing.assert_allclose(np.array([c['loc'] for c in check.info['chs']]),np.array([raw.info['chs'][raw.ch_names.index(name)]['loc'] for name in EEG19]),rtol=0,atol=1e-7,equal_nan=True)
                metadata={'n_channels':19,'eeg19_status':'完成','channel_names':json.dumps(EEG19),'fif_bytes':sum(f.stat().st_size for f in newfiles)}
                raw.close();raw=None;picked.close();picked=None;check.close();check=None
                # 先暂存旧文件，发布并检查新文件；清单成功更新后才删除旧副本。
                for file in old:file.rename(backup/file.name);moved.append(file)
                for file in newfiles:file.rename(folder/file.name);published.append(folder/file.name)
                check=mne.io.read_raw_fif(published[0],preload=False,verbose='ERROR');assert check.ch_names==EEG19;check.close();check=None
                row.update(metadata,fif_files=json.dumps([str(f) for f in published],ensure_ascii=False),error='')
                row['validation']=str(row.get('validation',''))+'；19导筛选：首中末与边界样本完全一致，事件及坐标保留'
                checkpoint(manifest,records)
                # 清单更新后新文件已生效；清理失败不能回滚并删除有效新文件。
                try:
                    for file in moved:(backup/file.name).unlink()
                    backup.rmdir();newdir.rmdir();work.rmdir()
                except OSError as cleanup_error:
                    row['error']='筛选已完成，临时副本清理待处理：'+str(cleanup_error)
                    checkpoint(manifest,records)
                status='完成'
            except Exception as error:
                for reader in [raw,picked,check]:
                    if reader is not None:reader.close()
                for file in published:
                    assert file.parent==folder
                    if file.exists():file.unlink()
                for file in moved:
                    saved=backup/file.name
                    if saved.exists():saved.rename(file)
                records[code]=original;records[code]['eeg19_status']='失败';records[code]['error']=type(error).__name__+': '+str(error)
                checkpoint(manifest,records);failures.append([code,records[code]['error']]);status='失败，原FIF保留'
            line=json.dumps({'index':index,'total':len(selected),'code':code,'status':status,'error':records[code].get('error','')},ensure_ascii=False)
            log.write(line+'\n');log.flush();print(line,flush=True)
            if args.pilot and failures:raise RuntimeError('19导小样验证失败')
    if not any(staging.iterdir()):staging.rmdir()
    complete=sum(r.get('eeg19_status')=='完成' and int(r['n_channels'])==19 for r in records.values())
    summary={'requested_records':len(selected),'eeg19_complete':complete,'total_records':len(records),'failed_this_run':len(failures),'manifest':str(manifest)}
    if not args.pilot:
        sections=[('结果',[f'当前{len(records)}条记录中{complete}条已保留19导EEG。本次处理{len(selected)}条，失败{len(failures)}条。'],['项目','数量'],[['19导完成',complete],['本次失败',len(failures)]]),('通道',['、'.join(EEG19),'去除A1/A2和辅助通道，仅筛选已有FIF；未从EDF重新转换、未滤波或重参考。EDF保持完整通道且未改动。'],[],[]),('验证',['保留通道的首、中、末和间断边界附近样本完全一致，样本数和采样率保持不变。事件、原时间元信息、BAD/EDGE boundary及电极坐标保留。新文件通过验证后替换旧FIF，未保留额外一套FIF。'],[],[]),('失败记录',[],['编号','原因'],failures)]
        summary['report']=write_report(base/'reports','03_19导EEG筛选报告_'+stamp,'19导EEG筛选报告',sections)
    print('FINAL_SUMMARY '+json.dumps(summary,ensure_ascii=False),flush=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--master',type=Path);parser.add_argument('--inspection',type=Path);parser.add_argument('--base',type=Path,required=True);parser.add_argument('--pilot',action='store_true');parser.add_argument('--codes',nargs='+',help='仅补跑指定评分编号');parser.add_argument('--join-discontinuous',action='store_true',help='用户确认后拼接间断记录并标记边界');parser.add_argument('--eeg19-only',action='store_true',help='仅筛选现有FIF为19导，验证后替换');args=parser.parse_args()
    if args.eeg19_only:
        select_eeg19(args);return
    if not args.master or not args.inspection:parser.error('修复转换需要 --master 与 --inspection')
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d_%H%M%S')
    inspection={r['code']:r for r in map(json.loads,args.inspection.read_text(encoding='utf-8').splitlines())}
    wb=load_workbook(args.master,read_only=True,data_only=True);rows=list(wb['受试者主表'].iter_rows(values_only=True));wb.close();cols={v:i for i,v in enumerate(rows[0])}
    subjects=[{'code':str(r[cols['编号']]),'group':r[cols['组别']],'guid':r[cols['记录GUID']],'source_edf':str(r[cols['服务器EDF路径']])} for r in rows[1:]]
    if args.codes:
        selected=set(args.codes);subjects=[item for item in subjects if item['code'] in selected]
        assert {item['code'] for item in subjects}==selected,'指定编号未在主表中全部找到'
    if args.pilot and not args.join_discontinuous:
        chosen={min((s for s in subjects if inspection[s['code']]['n_channels']==n),key=lambda s:inspection[s['code']]['file_bytes'])['code'] for n in [35,36]}
        chosen.add(min((s for s in subjects if s['group']=='健康组'),key=lambda s:inspection[s['code']]['file_bytes'])['code']);subjects=[s for s in subjects if s['code'] in chosen]
    edfs=args.base/'repaired/edf';fifs=args.base/'repaired/fif';staging=args.base/'repaired/.staging'
    for folder in [edfs,fifs,staging]:folder.mkdir(parents=True,exist_ok=True)
    manifest=args.base/'tables/03_修复对应表.csv';records={}
    if manifest.exists():
        with manifest.open(encoding='utf-8-sig',newline='') as f:records={r['code']:r for r in csv.DictReader(f)}
    logfile=args.base/'logs'/('03_修复转换_'+stamp+'.jsonl')
    with logfile.open('x',encoding='utf-8') as log:
        log.write(json.dumps({'stage':'configuration','master':str(args.master),'inspection':str(args.inspection),'pilot':args.pilot,'codes':args.codes,'join_discontinuous':args.join_discontinuous,'montage':'standard_1020','fif_format':'single','filtering':False,'rereference':False,'start_time':stamp},ensure_ascii=False)+'\n')
        for index,subject in enumerate(subjects,1):
            code=subject['code'];source=Path(subject['source_edf']);d=inspection[code];stat=source.stat()
            assert stat.st_size==d['file_bytes'] and stat.st_mtime_ns==d['mtime_ns'],'源文件自检查后发生变化'
            if args.join_discontinuous and code in records and records[code]['status']=='已修复EDF+D，FIF方案待确认':
                item=records[code].copy();item['error']='';stage=staging/code;stage.mkdir(exist_ok=True)
                target=Path(item['repaired_edf']);edf_stat=target.stat()
                try:
                    gaps=json.loads(item['time_gaps']);assert gaps
                    saved,metadata=convert(target,stage,code,source,d,joined_gaps=gaps)
                    assert all(not(fifs/file.name).exists() for file in saved)
                    published=[]
                    for file in saved:file.rename(fifs/file.name);published.append(str(fifs/file.name))
                    loaded=mne.io.read_raw_fif(published[0],preload=False,verbose='ERROR');assert loaded.n_times==metadata['n_samples'];loaded.close()
                    item.update(metadata,fif_files=json.dumps(published,ensure_ascii=False),fif_bytes=sum(Path(f).stat().st_size for f in published),status='成功');stage.rmdir()
                except Exception as error:
                    item['error']=type(error).__name__+': '+str(error)
                assert (target.stat().st_size,target.stat().st_mtime_ns)==(edf_stat.st_size,edf_stat.st_mtime_ns)
                records[code]=item;checkpoint(manifest,records)
                line=json.dumps({'index':index,'total':len(subjects),'code':code,'status':item['status'],'error':item['error']},ensure_ascii=False)
                log.write(line+'\n');log.flush();print(line,flush=True)
                if args.pilot and item['status']!='成功':raise RuntimeError('间断拼接小样验证未通过')
                continue
            if code in records and records[code]['status'] in ['成功','已修复EDF+D，FIF方案待确认']:
                if records[code]['status']!='成功':
                    assert Path(records[code]['repaired_edf']).is_file()
                    print(json.dumps({'index':index,'code':code,'status':'保留已修复间断EDF，FIF待确认'},ensure_ascii=False),flush=True);continue
                old=records[code];assert int(old['source_bytes'])==stat.st_size and int(old['source_mtime_ns'])==stat.st_mtime_ns
                assert Path(old['repaired_edf']).is_file() and all(Path(f).is_file() for f in json.loads(old['fif_files']))
                print(json.dumps({'index':index,'code':code,'status':'沿用已验证输出'},ensure_ascii=False),flush=True);continue
            item={**subject,'source_bytes':stat.st_size,'source_mtime_ns':stat.st_mtime_ns,'status':'失败','error':''};stage=staging/code
            try:
                stage.mkdir(exist_ok=True)
                target=edfs/(code+'.edf');assert not target.exists(),'不覆盖已有修复EDF'
                temporary=stage/(code+'.edf')
                gaps=[]
                if temporary.exists() and temporary.stat().st_size==stat.st_size+80:
                    original=hashlib.sha256();verified=hashlib.sha256()
                    with source.open('rb') as f:
                        f.seek(d['header_candidate'])
                        while block:=f.read(8*1024*1024):original.update(block)
                    with temporary.open('rb') as f:
                        f.seek(256*(d['n_channels']+1))
                        while block:=f.read(8*1024*1024):verified.update(block)
                    assert verified.hexdigest()==original.hexdigest(),'已有临时EDF数据不同，不覆盖'
                    item['data_sha256']=verified.hexdigest()
                else:
                    if temporary.exists():
                        assert temporary.resolve().parent==stage.resolve() and stage.resolve().is_relative_to(staging.resolve())
                        temporary.unlink()  # 仅移除本脚本未完成的数据副本，不涉及原始文件。
                    item['data_sha256'],gaps=repair(source,temporary,d)
                item['timekeeping_continuous']=not gaps;item['time_gaps']=json.dumps(gaps)
                if gaps:
                    temporary.rename(target)
                    item.update(repaired_edf=str(target),status='已修复EDF+D，FIF方案待确认',error='存在真实采集间断，保留时间轴，未转连续FIF')
                    records[code]=item;checkpoint(manifest,records);stage.rmdir()
                    line=json.dumps({'index':index,'total':len(subjects),'code':code,'status':item['status'],'gaps':len(gaps)},ensure_ascii=False)
                    log.write(line+'\n');log.flush();print(line,flush=True)
                    continue
                saved,metadata=convert(temporary,stage,code,source,d)
                for file in saved:assert not (fifs/file.name).exists()
                temporary.rename(target)
                published=[]
                for file in saved:file.rename(fifs/file.name);published.append(str(fifs/file.name))
                # 分片移动后再确认首文件可找到其余分片。
                loaded=mne.io.read_raw_fif(published[0],preload=False,verbose='ERROR');assert loaded.n_times==metadata['n_samples'];loaded.close()
                item.update(metadata,repaired_edf=str(target),fif_files=json.dumps(published,ensure_ascii=False),fif_bytes=sum(Path(f).stat().st_size for f in published),status='成功');stage.rmdir()
            except Exception as error:item['error']=type(error).__name__+': '+str(error)
            after=source.stat();assert (after.st_size,after.st_mtime_ns)==(stat.st_size,stat.st_mtime_ns)
            records[code]=item;checkpoint(manifest,records)
            line=json.dumps({'index':index,'total':len(subjects),'code':code,'status':item['status'],'error':item['error']},ensure_ascii=False);log.write(line+'\n');log.flush();print(line,flush=True)
            if args.pilot and item['status']!='成功':raise RuntimeError('小样验证失败，未开始批量转换')
    if not any(staging.iterdir()):staging.rmdir()
    success=[r for r in records.values() if r['status']=='成功'];pending=[r for r in records.values() if r['status']=='已修复EDF+D，FIF方案待确认'];failed=[r for r in records.values() if r['status'] not in ['成功','已修复EDF+D，FIF方案待确认']]
    summary={'success':len(success),'failed':len(failed),'pending_discontinuous':len(pending),'requested_codes':args.codes,'join_discontinuous':args.join_discontinuous,'manifest':str(manifest),'pilot':args.pilot,'log':str(logfile)}
    if not args.pilot:
        sections=[('结果',[f'当前主表{len(rows)-1}条记录，本次处理{len(subjects)}条。累计{len(success)}条完成EDF修复与FIF转换，{len(pending)}条已修复EDF+D但FIF待确认，失败{len(failed)}条。原始EDF未改动。'],['项目','数量'],[['完成EDF与FIF',len(success)],['已修复EDF+D，FIF待确认',len(pending)],['失败',len(failed)]]),('修复方法',['补齐缺失的80字节recording identification，未知采集者/设备字段用X；保留原患者字段、采集日期时间和数据区。记录数按文件长度填写。','按用户确认的10–20模板将第4、5通道恢复F4、C3，保持通道顺序与负增益。注释存在时逐记录核对时间标记连续性，保留事件文字和时长。中文事件在FIF中由MNE以XML字符实体存储，解码后逐项核对；显示时可用html.unescape还原。'],[],[]),('FIF转换',['EEG使用standard_1020标准坐标；ROC/LOC为眼电，EKG为心电，EMG为肌电，其余辅助通道保留为misc。未滤波、未重参考。用户确认的间断记录压缩暂停并拼接，连接处保留BAD/EDGE boundary；切窗和滤波不得跨边界。原始间断时长保留在对应表及事件extras，EDF+D不改动。','FIF使用MNE默认float32存储；事件起止时间也按FIF格式舍入，并在对应表记录最大偏差，原始精确事件字节保留在EDF。逐样验证预期舍入；大文件按FIF格式限制自动分片，分片属于同一记录。'],[],[]),('验证',['数据区完整SHA256一致。抽取首、中、末记录，由源数字值独立核对单位标定、极性和通道顺序；核对FIF样本数、通道类型、事件、采集时间和首中末样本。标准坐标是模板位置，不是个体实测坐标。'],[],[]),('产物',[str(edfs),str(fifs),str(manifest),str(logfile)],[],[]),('失败项',[],['编号','原因'],[[r['code'],r['error']] for r in failed])]
        summary['report']=write_report(args.base/'reports','03_EDF修复与FIF转换报告_'+stamp,'EDF修复与FIF转换报告',sections)
    print('FINAL_SUMMARY '+json.dumps(summary,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
